"""
テスト項目書 xlsx ヘルパー（汎用版）

使い方:
    import sys
    sys.path.insert(0, '<PROJECT_ROOT>/.claude/skills/test-spec-generator/scripts')
    from helpers import create_or_replace_sheet, update_row, verify_sheet_integrity, safe_unmerge

提供関数:
    - create_or_replace_sheet: 新シート作成 or 全体再生成
    - update_row: 1項目ピンポイント書き換え
    - verify_sheet_integrity: データ消失検出
    - safe_unmerge: スナップショット付きの安全な merge 解除
"""

from copy import copy
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


# ===== スタイル定義 =====

TITLE_FONT = Font(bold=True, size=14)
SECTION_FONT = Font(bold=True, size=12, color='FFFFFF')
SECTION_FILL = PatternFill('solid', fgColor='4472C4')
SUMMARY_TITLE_FONT = Font(bold=True, size=12, color='FFFFFF')
SUMMARY_TITLE_FILL = PatternFill('solid', fgColor='4472C4')
HEADER_FONT = Font(bold=True, size=11)
HEADER_FILL = PatternFill('solid', fgColor='D9E1F2')
LABEL_FILL = PatternFill('solid', fgColor='D9E1F2')
PREP_FONT = Font(size=11, color='000000')
PREP_FILL = PatternFill('solid', fgColor='F2F2F2')
NOTE_FONT = Font(italic=True, size=10, color='595959')
THIN = Side(border_style='thin', color='999999')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP_ALIGN = Alignment(wrap_text=True, vertical='top')
CENTER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT_ALIGN = Alignment(horizontal='left', vertical='center')

DEFAULT_TITLE = 'テスト項目書'
DEFAULT_LEGEND = '必要環境凡例: 開発=開発環境で実施可能 / データ要=テストデータ準備が必要 / 本番=本番環境でしか実施不可'

COLUMN_WIDTHS = {'A': 5, 'B': 12, 'C': 26, 'D': 50, 'E': 50, 'F': 8, 'G': 10, 'H': 24}

DATA_ROW_HEIGHT = 60
SECTION_ROW_HEIGHT = 22
PREP_SHORT_HEIGHT = 22
PREP_LONG_HEIGHT = 38
PREP_LONG_THRESHOLD = 100


# ===== 内部ヘルパー =====

def _setup_top_block(ws, title, scope, preconditions, legend=DEFAULT_LEGEND):
    """シート上部（タイトル + サマリ + 列ヘッダ）を r1〜r12 に配置"""

    # タイトルブロック r1〜r4
    ws['A1'] = title
    ws['A1'].font = TITLE_FONT
    ws.merge_cells('A1:H1')

    ws['A2'] = scope
    ws.merge_cells('A2:H2')

    ws['A3'] = preconditions
    ws['A3'].font = NOTE_FONT
    ws.merge_cells('A3:H3')

    ws['A4'] = legend
    ws['A4'].font = NOTE_FONT
    ws.merge_cells('A4:H4')

    # r5: 空行
    ws.row_dimensions[5].height = 8

    # r6〜r10: 結果サマリ（B6:C10、右に1列ずらした位置）
    ws['B6'] = '結果サマリ'
    ws['B6'].font = SUMMARY_TITLE_FONT
    ws['B6'].fill = SUMMARY_TITLE_FILL
    ws['B6'].alignment = Alignment(horizontal='center', vertical='center')
    ws.merge_cells('B6:C6')
    ws.row_dimensions[6].height = 22

    summary = [
        (7, '全項目数', '=COUNT($A$15:$A$200)'),
        (8, '合格', '=COUNTIF($F$15:$F$200,"o")'),
        (9, '不合格', '=COUNTIF($F$15:$F$200,"x")+COUNTIF($F$15:$F$200,"×")'),
        (10, '未実施', '=C7-C8-C9'),
    ]
    for r, label, formula in summary:
        ws.cell(row=r, column=2, value=label).font = HEADER_FONT
        ws.cell(row=r, column=2).fill = LABEL_FILL
        ws.cell(row=r, column=2).alignment = LEFT_ALIGN
        ws.cell(row=r, column=2).border = BORDER
        ws.cell(row=r, column=3, value=formula).alignment = CENTER_ALIGN
        ws.cell(row=r, column=3).border = BORDER
        ws.row_dimensions[r].height = 20

    # r11: 区切り空行（小）
    ws.row_dimensions[11].height = 8

    # r12: 列ヘッダ
    headers = ['#', 'カテゴリ', '機能・項目', '操作手順', '期待結果', '結果', '必要環境', '備考']
    for col_idx, h in enumerate(headers, 1):
        c = ws.cell(row=12, column=col_idx, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = CENTER_ALIGN
        c.border = BORDER
    ws.row_dimensions[12].height = 22

    # 列幅
    for col, w in COLUMN_WIDTHS.items():
        ws.column_dimensions[col].width = w


def _write_section(ws, start_row, sec, item_no_start):
    """1セクション（ヘッダ + 事前準備 + 各項目）を書き込み、終了行と次の item_no を返す"""

    row = start_row

    # セクションヘッダ
    ws.cell(row=row, column=1, value=sec['title'])
    ws.cell(row=row, column=1).font = SECTION_FONT
    ws.cell(row=row, column=1).fill = SECTION_FILL
    ws.cell(row=row, column=1).alignment = Alignment(vertical='center')
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
    ws.row_dimensions[row].height = SECTION_ROW_HEIGHT
    row += 1

    # 事前準備行
    if sec.get('prep'):
        prep_text = sec['prep']
        ws.cell(row=row, column=1, value=prep_text)
        ws.cell(row=row, column=1).font = PREP_FONT
        ws.cell(row=row, column=1).fill = PREP_FILL
        ws.cell(row=row, column=1).alignment = Alignment(wrap_text=True, vertical='center')
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        if len(prep_text) > PREP_LONG_THRESHOLD:
            ws.row_dimensions[row].height = PREP_LONG_HEIGHT
        else:
            ws.row_dimensions[row].height = PREP_SHORT_HEIGHT
        row += 1

    # 各項目
    item_no = item_no_start
    for item in sec['items']:
        cat, feat, op, expected, env, remark = item
        ws.cell(row=row, column=1, value=item_no).alignment = CENTER_ALIGN
        ws.cell(row=row, column=1).border = BORDER
        ws.cell(row=row, column=2, value=cat).alignment = WRAP_ALIGN
        ws.cell(row=row, column=2).border = BORDER
        ws.cell(row=row, column=3, value=feat).alignment = WRAP_ALIGN
        ws.cell(row=row, column=3).border = BORDER
        ws.cell(row=row, column=4, value=op).alignment = WRAP_ALIGN
        ws.cell(row=row, column=4).border = BORDER
        ws.cell(row=row, column=5, value=expected).alignment = WRAP_ALIGN
        ws.cell(row=row, column=5).border = BORDER
        ws.cell(row=row, column=6).alignment = CENTER_ALIGN
        ws.cell(row=row, column=6).border = BORDER
        ws.cell(row=row, column=7, value=env).alignment = CENTER_ALIGN
        ws.cell(row=row, column=7).border = BORDER
        ws.cell(row=row, column=8, value=remark).alignment = WRAP_ALIGN
        ws.cell(row=row, column=8).border = BORDER
        ws.row_dimensions[row].height = DATA_ROW_HEIGHT
        row += 1
        item_no += 1

    return row, item_no


def _collect_user_input(ws):
    """既存シートから F列（結果）と H列（備考）を {item_no: (result, remark)} で収集"""
    user_input = {}
    for r in range(13, ws.max_row + 1):
        a = ws.cell(row=r, column=1).value
        if not isinstance(a, int):
            continue
        result = ws.cell(row=r, column=6).value
        remark = ws.cell(row=r, column=8).value
        user_input[a] = (result, remark)
    return user_input


def _apply_user_input(ws, user_input):
    """収集した F・H列値を、新しいシート構造に書き戻す"""
    for r in range(13, ws.max_row + 1):
        a = ws.cell(row=r, column=1).value
        if not isinstance(a, int):
            continue
        if a in user_input:
            result, remark = user_input[a]
            if result is not None:
                ws.cell(row=r, column=6).value = result
                ws.cell(row=r, column=6).alignment = CENTER_ALIGN
            if remark is not None:
                ws.cell(row=r, column=8).value = remark
                ws.cell(row=r, column=8).alignment = WRAP_ALIGN


# ===== Public API =====

def create_or_replace_sheet(
    xlsx_path,
    sheet_name,
    sections,
    title=DEFAULT_TITLE,
    scope='',
    preconditions='',
    legend=DEFAULT_LEGEND,
    preserve_user_input=True,
):
    """
    指定シートを新規作成 or 全体再生成する。

    Args:
        xlsx_path: 出力先 xlsx のパス。存在しなければ新規作成
        sheet_name: シート名（既存があれば置き換え）
        sections: list of dict
            [
                {
                    'title': '【セクション1】〇〇',
                    'prep': '事前準備: ...' (or None),
                    'items': [
                        (cat, feat, op, expected, env, remark),
                        ...
                    ],
                },
                ...
            ]
        title: シート上部のタイトル
        scope: 対象機能の説明
        preconditions: 前提条件
        legend: 必要環境凡例
        preserve_user_input: True の場合、既存シートの F・H列値を保持する

    Returns:
        生成したシート（worksheet オブジェクト）
    """
    import os
    if os.path.exists(xlsx_path):
        wb = load_workbook(xlsx_path)
    else:
        wb = Workbook()
        if 'Sheet' in wb.sheetnames:
            del wb['Sheet']

    user_input = {}
    if preserve_user_input and sheet_name in wb.sheetnames:
        ws_old = wb[sheet_name]
        user_input = _collect_user_input(ws_old)

    if sheet_name in wb.sheetnames:
        del wb[sheet_name]

    ws = wb.create_sheet(sheet_name, 0)

    _setup_top_block(ws, title, scope, preconditions, legend)

    row = 13
    item_no = 1
    for sec in sections:
        row, item_no = _write_section(ws, row, sec, item_no)

    wb.save(xlsx_path)

    if user_input:
        wb2 = load_workbook(xlsx_path)
        ws2 = wb2[sheet_name]
        _apply_user_input(ws2, user_input)
        wb2.save(xlsx_path)

    return ws


def update_row(xlsx_path, sheet_name, item_no, updates):
    """
    指定項目（item_no）の行を部分更新する。

    Args:
        xlsx_path: xlsx のパス
        sheet_name: シート名
        item_no: 項目番号（A 列の整数値）
        updates: dict
            キー: 'category' / 'feature' / 'operation' / 'expected' / 'result' / 'env' / 'remark'
            値: 書き換える文字列

    Returns:
        更新した行番号（見つからない場合は None）
    """
    field_to_col = {
        'category': 2,
        'feature': 3,
        'operation': 4,
        'expected': 5,
        'result': 6,
        'env': 7,
        'remark': 8,
    }

    wb = load_workbook(xlsx_path)
    ws = wb[sheet_name]

    target_row = None
    for r in range(13, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == item_no:
            target_row = r
            break

    if target_row is None:
        return None

    for field, value in updates.items():
        col = field_to_col.get(field)
        if col is None:
            raise ValueError(f'Unknown field: {field}. Allowed: {list(field_to_col)}')
        cell = ws.cell(row=target_row, column=col)
        cell.value = value
        if col in (6, 7):
            cell.alignment = CENTER_ALIGN
        else:
            cell.alignment = WRAP_ALIGN

    wb.save(xlsx_path)
    return target_row


def verify_sheet_integrity(xlsx_path, sheet_name):
    """
    シートの整合性を検証する。データ消失を検出。

    Returns:
        list of dict: 検出した問題のリスト
            [{'row': N, 'item_no': X, 'issue': 'B列が空'}, ...]
    """
    wb = load_workbook(xlsx_path)
    ws = wb[sheet_name]

    issues = []

    for r in range(13, ws.max_row + 1):
        a = ws.cell(row=r, column=1).value
        if not isinstance(a, int):
            continue
        b = ws.cell(row=r, column=2).value
        if b is None:
            issues.append({
                'row': r,
                'item_no': a,
                'issue': 'B列（カテゴリ）が空。データ消失の可能性あり',
            })

    section_prep_rows = []
    for r in range(13, ws.max_row + 1):
        a = ws.cell(row=r, column=1).value
        if not isinstance(a, str):
            continue
        if a.startswith('【') or a.startswith('事前準備'):
            section_prep_rows.append((r, a))

    for r, a in section_prep_rows:
        is_merged = any(
            mr.min_row == r and mr.max_row == r and mr.max_col == 8
            for mr in ws.merged_cells.ranges
        )
        if not is_merged:
            issues.append({
                'row': r,
                'item_no': None,
                'issue': f'セクション/事前準備行の A:H マージが未設定: {a[:40]}',
            })

    return issues


def safe_unmerge(ws, merge_range_str):
    """
    merge を解除する前にスナップショットを取り、解除後に値を書き戻す。

    Args:
        ws: worksheet
        merge_range_str: 'A7:H7' など
    """
    from openpyxl.utils import range_boundaries
    min_col, min_row, max_col, max_row = range_boundaries(merge_range_str)

    snapshot = {}
    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            val = ws.cell(row=r, column=c).value
            if val is not None:
                snapshot[(r, c)] = val

    ws.unmerge_cells(merge_range_str)

    for (r, c), val in snapshot.items():
        ws.cell(row=r, column=c).value = val
